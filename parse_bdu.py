import requests
import urllib3
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

url_1 = 'https://bdu.fstec.ru/threat/ubi.00'
url_2 = 'https://bdu.fstec.ru/threat/ubi.0'
url_3 = 'https://bdu.fstec.ru/threat/ubi.'

headers = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/75.0.3770.142 Safari/537.36'
}

ubi_list = []
description_list = []
intruder_list = []
influence_object_list = []
conclusion_list = []


for i in range(1, 221):
    if i < 10:
        response = requests.get(url_1+str(i), verify=False, headers=headers)
        soup = BeautifulSoup(response.text, 'html.parser')
        ubi_number = soup.find('div', {"class": "col-sm-11"}).text.replace('\n', '').split(':')[0]
        description = soup.find('div', {"class": "col-sm-11"}).text.replace('\n', '').split(':')[1]
        intruder = soup.find_all('tr')[1].text.replace('\n', '').replace('  ','').replace('ом', 'ом\n')
        influence_object = soup.find_all('tr')[2].text.replace('\n', '').replace('  ','')
        conclusion = soup.find_all('tr')[3].text.replace('\n', '').replace('  ', '').replace('ти', 'ти\n')
        print(f"{ubi_number}")
        print(f"Описание: {description}")
        print(f"Нарушители:\n{intruder.replace('Источники угрозы', '')[:-1]}")
        print(f"Объект воздействия: {influence_object.replace('Объект воздействия', '')}")
        print(f"Последствия реализации угрозы:\n{conclusion.replace('Последствия реализации угрозы', '')}")
        print('---------------------------------------------------------------------------------------------------')
        ubi_list.append(ubi_number)
        description_list.append(description)
        intruder_list.append(intruder.replace('Источники угрозы', '')[:-1])
        influence_object_list.append(influence_object.replace('Объект воздействия', ''))
        conclusion_list.append(conclusion.replace('Последствия реализации угрозы', ''))
    elif i < 100:
        response = requests.get(url_2 + str(i), verify=False, headers=headers)
        soup = BeautifulSoup(response.text, 'html.parser')
        ubi_number = soup.find('div', {"class": "col-sm-11"}).text.replace('\n', '').split(':')[0]
        description = soup.find('div', {"class": "col-sm-11"}).text.replace('\n', '').split(':')[1]
        intruder = soup.find_all('tr')[1].text.replace('\n', '').replace('  ', '').replace('ом', 'ом\n')
        influence_object = soup.find_all('tr')[2].text.replace('\n', '').replace('  ', '')
        conclusion = soup.find_all('tr')[3].text.replace('\n', '').replace('  ', '').replace('ти', 'ти\n')
        print(f"{ubi_number}")
        print(f"Описание: {description}")
        print(f"Нарушители:\n{intruder.replace('Источники угрозы', '')[:-1]}")
        print(f"Объект воздействия: {influence_object.replace('Объект воздействия', '')}")
        print(f"Последствия реализации угрозы:\n{conclusion.replace('Последствия реализации угрозы', '')}")
        print('---------------------------------------------------------------------------------------------------')
        ubi_list.append(ubi_number)
        description_list.append(description)
        intruder_list.append(intruder.replace('Источники угрозы', '')[:-1])
        influence_object_list.append(influence_object.replace('Объект воздействия', ''))
        conclusion_list.append(conclusion.replace('Последствия реализации угрозы', ''))
    else:
        response = requests.get(url_3 + str(i), verify=False, headers=headers)
        soup = BeautifulSoup(response.text, 'html.parser')
        try:
            ubi_number = soup.find('div', {"class": "col-sm-11"}).text.replace('\n', '').split(':')[0]
            description = soup.find('div', {"class": "col-sm-11"}).text.replace('\n', '').split(':')[1]
            intruder = soup.find_all('tr')[1].text.replace('\n', '').replace('  ', '').replace('ом', 'ом\n')
            influence_object = soup.find_all('tr')[2].text.replace('\n', '').replace('  ', '')
            conclusion = soup.find_all('tr')[3].text.replace('\n', '').replace('  ', '').replace('ти', 'ти\n')
            print(f"{ubi_number}")
            print(f"Описание: {description}")
            print(f"Нарушители:\n{intruder.replace('Источники угрозы', '')[:-1]}")
            print(f"Объект воздействия: {influence_object.replace('Объект воздействия', '')}")
            print(f"Последствия реализации угрозы:\n{conclusion.replace('Последствия реализации угрозы', '')}")
            print('---------------------------------------------------------------------------------------------------')
            ubi_list.append(ubi_number)
            description_list.append(description)
            intruder_list.append(intruder.replace('Источники угрозы', '')[:-1])
            influence_object_list.append(influence_object.replace('Объект воздействия', ''))
            conclusion_list.append(conclusion.replace('Последствия реализации угрозы', ''))
        except IndexError:
            ubi_number = soup.find('div', {"class": "col-sm-11"}).text.replace('\n', '').split(':')[0]
            description = soup.find('div', {"class": "col-sm-11"}).text.replace('\n', '').split(':')[1]
            intruder = soup.find_all('tr')[1].text.replace('\n', '').replace('  ', '').replace('ом', 'ом\n')
            influence_object = soup.find_all('tr')[2].text.replace('\n', '').replace('  ', '')
            print(f"{ubi_number}")
            print(f"Описание: {description}")
            print(f"Объект воздействия: {intruder.replace('Объект воздействия', '')}")
            print(f"Последствия реализации угрозы:\n{influence_object.replace('Последствия реализации угрозы', '')}")
            print('---------------------------------------------------------------------------------------------------')
            ubi_list.append(ubi_number)
            description_list.append(description)
            intruder_list.append('Нет')
            influence_object_list.append(intruder.replace('Объект воздействия', ''))
            conclusion_list.append(influence_object.replace('Последствия реализации угрозы', ''))

def create_word(ubi_list, description_list, intruder_list, influence_object_list, conclusion_list):
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    table = doc.add_table(rows=221, cols=5, style='Table Grid')
    for i in range(1, 221):
        table.cell(i,0).text = (ubi_list[i-1])
        table.cell(i, 1).text = (description_list[i-1])
        table.cell(i, 2).text = (intruder_list[i-1])
        table.cell(i, 3).text = (influence_object_list[i-1])
        table.cell(i, 4).text = (conclusion_list[i-1])
    table.cell(0, 0).text = '№ УБИ'
    table.cell(0, 1).text = 'Описание'
    table.cell(0, 2).text = 'Источники угрозы'
    table.cell(0, 3).text = 'Объект воздействия'
    table.cell(0, 4).text = 'Последствия реализации угрозы'
    doc.save('Угрозы безопасности информации.docx')

create_word(ubi_list, description_list, intruder_list, influence_object_list, conclusion_list)